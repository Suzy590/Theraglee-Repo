"""Build the spreadsheet from rows/therapists.json plus rows/emails.json."""
import json, os, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUT_DIR = sys.argv[1]; OUT = sys.argv[2]
rows = json.load(open(os.path.join(OUT_DIR, "rows", "therapists.json")))
ep = os.path.join(OUT_DIR, "rows", "emails.json")
emails = json.load(open(ep)) if os.path.exists(ep) else {}
looked = set()
import glob
for fn in glob.glob(os.path.join(OUT_DIR, "sites", "*-rows.json")):
    for it in json.load(open(fn)):
        looked.add(it["npi"])

FONT = "Arial"
HDR = ["First name", "Last name", "Credential (self-reported)", "License type", "Other license types",
       "License number(s) on file", "City", "ZIP", "Practice address", "Phone", "Email", "Website",
       "Email found on", "Private practice (sole proprietor)", "Website looked up?", "County", "NPI", "NPI last updated", "Note"]
WIDTHS = [14, 16, 14, 30, 28, 22, 16, 7, 34, 14, 30, 34, 34, 12, 11, 14, 12, 12, 40]

def sheet_for(wb, name, data):
    ws = wb.create_sheet(name)
    ws.append(HDR)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F6B4F")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in data:
        e = emails.get(r["npi"], {})
        ws.append([v if v != "" else None for v in [r["first_name"], r["last_name"], r["credential"], r["primary_license"],
                   r["all_licenses"] if ";" in r["all_licenses"] else "", r["license_numbers"], r["city"], r["zip"],
                   r["address"], r["phone"], e.get("email", ""), e.get("website", ""), e.get("email_source", ""),
                   "Yes" if r["sole_proprietor"] == "YES" else "No",
                   "Yes" if r["npi"] in looked else "Not yet", r["county"], r["npi"], r["npi_updated"], r["flag"]]])
    for i, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=FONT, size=10)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HDR))}{ws.max_row}"
    return ws

wb = Workbook()
summary = wb.active; summary.title = "Summary"
order = ["Ventura", "Santa Barbara", "Los Angeles"]
by = {c: sorted([r for r in rows if r["county"] == c], key=lambda r: (r["sole_proprietor"] != "YES", r["last_name"], r["first_name"])) for c in order}
for c in order:
    sheet_for(wb, c, by[c])

# Summary: counts are live formulas over the county tabs.
s = summary
s["A1"] = "Licensed therapists in Ventura, Santa Barbara and Los Angeles counties"; s["A1"].font = Font(name=FONT, bold=True, size=14)
s["A2"] = f"Built {datetime.date.today():%B %d, %Y} from the NPI Registry (npiregistry.cms.hhs.gov), the US government's public list of healthcare providers. Each county tab is one row per provider."
s["A2"].font = Font(name=FONT, size=10, italic=True)
s["A4"] = "County"; s["B4"] = "Providers"; s["C4"] = "Private practice"; s["D4"] = "Websites looked up"; s["E4"] = "Emails found"
for c in s[4]: c.font = Font(name=FONT, bold=True)
for i, c in enumerate(order, start=5):
    q = f"'{c}'"
    s[f"A{i}"] = c
    s[f"B{i}"] = f"=COUNTA({q}!Q:Q)-1"
    s[f"C{i}"] = f'=COUNTIF({q}!N:N,"Yes")'
    s[f"D{i}"] = f'=COUNTIF({q}!O:O,"Yes")'
    s[f"E{i}"] = f"=COUNTA({q}!K:K)-1"
s["A8"] = "Total"; s["A8"].font = Font(name=FONT, bold=True)
for col in "BCDE":
    s[f"{col}8"] = f"=SUM({col}5:{col}7)"; s[f"{col}8"].font = Font(name=FONT, bold=True)
for r in range(5, 9):
    for col in "BCDE":
        s[f"{col}{r}"].number_format = "#,##0"

s["A10"] = "By license type"; s["A10"].font = Font(name=FONT, bold=True)
s["A11"] = "License type"; s["B11"] = "Ventura"; s["C11"] = "Santa Barbara"; s["D11"] = "Los Angeles"
for c in s[11]: c.font = Font(name=FONT, bold=True)
types = sorted({r["primary_license"] for r in rows})
for j, t in enumerate(types, start=12):
    s[f"A{j}"] = t
    for col, c in zip("BCD", order):
        s[f"{col}{j}"] = f'=COUNTIF(\'{c}\'!D:D,A{j})'; s[f"{col}{j}"].number_format = "#,##0"

n = 12 + len(types) + 1
notes = [
    "How to read this workbook",
    "• License type comes from the taxonomy code on the provider's NPI record, mapped to the license types on Theraglee's therapist sign-up form. 'Credential' is whatever the provider typed on their own NPI record (LMFT, PsyD, MSW...).",
    "• 'License number(s) on file' is what the provider reported to the NPI Registry. It is not a verification. Rows without one may be associates, interns or agency staff; check with the California BBS or Board of Psychology before relying on it.",
    "• 'Private practice' is the NPI sole-proprietor flag. 'Yes' is the best sign the person runs their own practice and would want client referrals.",
    "• Email and Website were found by searching for the provider's own practice website and reading its public contact page. Blank means no site was found or the site lists no address. 'Not yet' in 'Website looked up?' means that row has not been searched.",
    "• County is assigned from the practice address ZIP code (or city name when the ZIP is not in the public ZIP table). Westlake Village (91361) sits on the Ventura/Los Angeles line and is listed under Ventura.",
    "• Note flags a license on file from another state, or a provider whose only address in these counties is a mailing address.",
    "• Before emailing anyone on this list, follow the three legal rules in docs/outreach-email.md: an honest subject line, a working unsubscribe link, and a postal address in the footer.",
]
for k, line in enumerate(notes):
    cell = s[f"A{n + k}"]; cell.value = line
    cell.font = Font(name=FONT, bold=(k == 0), size=10); cell.alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells(f"A{n + k}:F{n + k}")
    if k: s.row_dimensions[n + k].height = 32
for col, w in zip("ABCDEF", (44, 14, 16, 18, 14, 14)):
    s.column_dimensions[col].width = w
for row in s.iter_rows():
    for c in row:
        if not c.font or c.font.name != FONT:
            c.font = Font(name=FONT, size=10, bold=c.font.bold if c.font else False)
wb.save(OUT)
print("saved", OUT, "rows:", {c: len(v) for c, v in by.items()}, "emails:", sum(1 for v in emails.values() if v.get("email")))
